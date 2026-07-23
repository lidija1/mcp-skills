from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    load_workbook = None


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MAX_FIELD_VALUES = 20
MAX_EXAMPLES = 8
EXCLUDED_DIRS = {
    ".git",
    ".idea",
    ".mypy_cache",
    ".pytest_cache",
    ".tox",
    ".venv",
    "__pycache__",
    "allure-report",
    "build",
    "dist",
    "node_modules",
    "target",
    "venv",
}


def _safe_lob(value: str) -> str:
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1-\2", value.strip())
    cleaned = re.sub(r"[^a-z0-9_-]+", "-", value.lower())
    return cleaned.strip("-_")


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="ignore")


def _read_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    return str(value).strip()


def _rel(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve())).replace("\\", "/")
    except ValueError:
        return str(path.resolve()).replace("\\", "/")


def _iter_files(root: Path, patterns: list[str]) -> list[Path]:
    files: list[Path] = []
    for pattern in patterns:
        files.extend(
            path for path in root.rglob(pattern)
            if path.is_file() and not any(part in EXCLUDED_DIRS for part in path.parts)
        )
    return sorted(set(files))


def _existing_roots(root: Path, candidates: list[str]) -> list[Path]:
    return [(root / candidate).resolve() for candidate in candidates if (root / candidate).exists()]


def _detect_framework_type(root: Path, requested: str) -> str:
    if requested != "auto":
        return requested
    if (root / "pom.xml").exists() or (root / "src/test/java").exists():
        return "java"
    if (
        (root / "pyproject.toml").exists()
        or (root / "pytest.ini").exists()
        or (root / "requirements.txt").exists()
        or (root / "tests").exists()
    ):
        return "python"
    return "generic"


def _infer_lob_from_path(path: Path, root: Path) -> str:
    parts = list(path.relative_to(root).parts)
    lowered = [part.lower() for part in parts]
    for marker in ("features", "pages", "steps", "tests", "test"):
        if marker in lowered:
            index = lowered.index(marker)
            if index + 1 < len(parts) - 1:
                return _safe_lob(parts[index + 1])
    return ""


def _parse_feature(path: Path, root: Path) -> dict[str, Any]:
    text = _read_text(path)
    lines = text.splitlines()
    feature_name = ""
    tags: list[str] = []
    scenarios: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    example_headers: list[str] = []
    in_examples = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("@"):
            tags = line.split()
            continue
        if line.lower().startswith("feature:"):
            feature_name = line.split(":", 1)[1].strip()
            continue
        if re.match(r"^(Scenario|Scenario Outline):", line, re.IGNORECASE):
            if current:
                scenarios.append(current)
            scenario_type, name = line.split(":", 1)
            current = {
                "name": name.strip(),
                "type": scenario_type.strip(),
                "tags": tags,
                "steps": [],
                "examples": [],
            }
            tags = []
            example_headers = []
            in_examples = False
            continue
        if current is None:
            continue
        if line.lower().startswith("examples:"):
            in_examples = True
            example_headers = []
            continue
        if in_examples and line.startswith("|") and line.endswith("|"):
            values = [part.strip() for part in line.strip("|").split("|")]
            if not example_headers:
                example_headers = values
            else:
                current["examples"].append(dict(zip(example_headers, values)))
            continue
        if re.match(r"^(Given|When|Then|And|But)\b", line):
            current["steps"].append(line)
            in_examples = False

    if current:
        scenarios.append(current)

    excel_refs: list[dict[str, str]] = []
    for scenario in scenarios:
        for row in scenario["examples"]:
            if "ExcelData" in row or "SHEET" in row or "TC_ID" in row:
                excel_refs.append(
                    {
                        "feature": _rel(path, root),
                        "scenario": scenario["name"],
                        "excel_data": row.get("ExcelData", ""),
                        "sheet": row.get("SHEET", ""),
                        "tc_id": row.get("TC_ID", ""),
                    }
                )

    return {
        "path": _rel(path, root),
        "lob": _infer_lob_from_path(path, root),
        "feature": feature_name,
        "scenario_count": len(scenarios),
        "scenarios": scenarios,
        "excel_references": excel_refs,
    }


def _parse_java(path: Path, root: Path) -> dict[str, Any]:
    text = _read_text(path)
    class_match = re.search(r"\bclass\s+([A-Za-z0-9_]+)", text)
    annotations = re.findall(r'@(Given|When|Then|And|But)\("([^"]+)"\)', text)
    data_fields = sorted(set(re.findall(r'data\.get\("([^"]+)"\)', text)))
    methods = re.findall(
        r"\b(?:public|private|protected)\s+[A-Za-z0-9_<>, ?\[\]]+\s+([A-Za-z0-9_]+)\s*\(",
        text,
    )
    return {
        "path": _rel(path, root),
        "language": "java",
        "lob": _infer_lob_from_path(path, root),
        "class": class_match.group(1) if class_match else path.stem,
        "step_definitions": [{"keyword": key, "text": value} for key, value in annotations],
        "data_fields": data_fields,
        "method_names": sorted(set(methods))[:80],
    }


def _parse_python(path: Path, root: Path) -> dict[str, Any]:
    text = _read_text(path)
    classes = re.findall(r"^\s*class\s+([A-Za-z_][A-Za-z0-9_]*)", text, re.MULTILINE)
    functions = re.findall(r"^\s*def\s+([A-Za-z_][A-Za-z0-9_]*)\s*\(", text, re.MULTILINE)
    step_decorators = re.findall(
        r"@(given|when|then|step)\(\s*[rRuUbBfF]*[\"']([^\"']+)[\"']",
        text,
        re.IGNORECASE,
    )
    data_fields = set()
    patterns = [
        r'\bdata\.get\(\s*[rRuUbBfF]*["\']([^"\']+)["\']',
        r'\bcontext\.config\.userdata\.get\(\s*[rRuUbBfF]*["\']([^"\']+)["\']',
        r'\b(?:row|record|payload|persona|test_data|data)\s*\[\s*[rRuUbBfF]*["\']([^"\']+)["\']\s*\]',
        r'\b(?:os\.)?getenv\(\s*[rRuUbBfF]*["\']([^"\']+)["\']',
    ]
    for pattern in patterns:
        data_fields.update(re.findall(pattern, text))

    return {
        "path": _rel(path, root),
        "language": "python",
        "lob": _infer_lob_from_path(path, root),
        "class": classes[0] if classes else path.stem,
        "classes": sorted(set(classes))[:40],
        "step_definitions": [{"keyword": key.lower(), "text": value} for key, value in step_decorators],
        "data_fields": sorted(data_fields),
        "method_names": sorted(set(functions))[:80],
    }


def _parse_code(path: Path, root: Path) -> dict[str, Any]:
    if path.suffix.lower() == ".py":
        return _parse_python(path, root)
    return _parse_java(path, root)


def _find_header_row(sheet, max_scan_rows: int = 10) -> int:
    best_row = 1
    best_score = 0
    for row_idx, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_scan_rows), values_only=True),
        start=1,
    ):
        values = [_read_cell(value) for value in row]
        non_empty = [value for value in values if value]
        score = len(non_empty) + len(set(non_empty))
        if score > best_score:
            best_row = row_idx
            best_score = score
    return best_row


def _headers_for_sheet(sheet, header_row: int) -> list[tuple[int, str]]:
    headers = []
    seen: Counter[str] = Counter()
    rows = sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
    header_values = next(rows, ())
    for column, raw_value in enumerate(header_values, start=1):
        header = _read_cell(raw_value)
        if not header:
            continue
        seen[header] += 1
        if seen[header] > 1:
            header = f"{header}_{seen[header]}"
        headers.append((column, header))
    return headers


def _scan_workbook(path: Path, root: Path, max_rows: int) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("Python dependency 'openpyxl' is missing. Install requirements.txt to scan Excel files.")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets: list[dict[str, Any]] = []

        for sheet in workbook.worksheets:
            header_row = _find_header_row(sheet)
            headers = _headers_for_sheet(sheet, header_row)
            header_by_column = {column: header for column, header in headers}
            values_by_field: dict[str, list[str]] = defaultdict(list)
            tc_ids: list[str] = []
            rows_scanned = 0
            sample_max_row = min(sheet.max_row, header_row + max_rows)

            for row_values_raw in sheet.iter_rows(
                min_row=header_row + 1,
                max_row=sample_max_row,
                values_only=True,
            ):
                has_value = False
                row_values: dict[str, str] = {}
                for column, raw_value in enumerate(row_values_raw, start=1):
                    header = header_by_column.get(column)
                    if not header:
                        continue
                    value = _read_cell(raw_value)
                    row_values[header] = value
                    if value:
                        has_value = True
                    values_by_field[header].append(value)
                if not has_value:
                    continue
                rows_scanned += 1
                first_header = headers[0][1] if headers else ""
                candidate = row_values.get(first_header, "")
                if candidate and re.search(r"(TC_|TEST|ID)", candidate, re.IGNORECASE):
                    tc_ids.append(candidate)

            field_summary = []
            for field, values in values_by_field.items():
                non_empty = [value for value in values if value]
                unique = sorted(set(non_empty))
                item = {
                    "name": field,
                    "populated_count": len(non_empty),
                    "unique_count": len(unique),
                }
                if 0 < len(unique) <= MAX_FIELD_VALUES:
                    item["observed_values"] = unique[:MAX_FIELD_VALUES]
                elif unique:
                    item["examples"] = unique[:MAX_EXAMPLES]
                field_summary.append(item)

            sheets.append(
                {
                    "name": sheet.title,
                    "header_row": header_row,
                    "rows_scanned": rows_scanned,
                    "total_rows": max(sheet.max_row - header_row, 0),
                    "headers": [header for _, header in headers],
                    "tc_id_examples": tc_ids[:MAX_EXAMPLES],
                    "fields": field_summary,
                }
            )

        return {
            "path": _rel(path, root),
            "sheets": sheets,
        }
    finally:
        workbook.close()


def _build_lob_summary(features: list[dict[str, Any]], code_files: list[dict[str, Any]]) -> dict[str, Any]:
    lobs: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "feature_files": [],
            "code_files": [],
            "page_classes": [],
            "step_classes": [],
            "scenario_count": 0,
            "data_fields": set(),
            "test_data_references": [],
        }
    )

    for feature in features:
        lob = feature.get("lob") or "unknown"
        item = lobs[lob]
        item["feature_files"].append(feature["path"])
        item["scenario_count"] += feature["scenario_count"]
        item["test_data_references"].extend(feature["excel_references"])

    for code_file in code_files:
        lob = code_file.get("lob") or "shared"
        item = lobs[lob]
        item["code_files"].append(code_file["path"])
        if code_file["step_definitions"]:
            item["step_classes"].append(code_file["path"])
        else:
            item["page_classes"].append(code_file["path"])
        item["data_fields"].update(code_file["data_fields"])

    normalized: dict[str, Any] = {}
    for lob, item in sorted(lobs.items()):
        normalized[lob] = {
            "feature_files": sorted(set(item["feature_files"])),
            "code_files": sorted(set(item["code_files"])),
            "page_classes": sorted(set(item["page_classes"])),
            "step_classes": sorted(set(item["step_classes"])),
            "scenario_count": item["scenario_count"],
            "data_fields": sorted(item["data_fields"]),
            "test_data_references": item["test_data_references"],
        }
    return normalized


def build_knowledge(args: argparse.Namespace) -> dict[str, Any]:
    root = Path(args.root).resolve()
    framework_type = _detect_framework_type(root, args.framework_type)
    feature_roots = _existing_roots(root, [args.features]) if args.features else []
    code_roots = _existing_roots(root, [args.code or args.java]) if (args.code or args.java) else []
    test_data_roots = _existing_roots(root, [args.test_data]) if args.test_data else []
    lob_filter = _safe_lob(args.lob) if args.lob else ""

    if not feature_roots:
        feature_roots = _existing_roots(
            root,
            ["src/test/java/features", "src/test/resources/features", "features", "tests/features"],
        )
    if not code_roots:
        if framework_type == "java":
            code_roots = _existing_roots(root, ["src/test/java", "src/main/java"])
        elif framework_type == "python":
            code_roots = _existing_roots(root, ["tests", "src", "pages", "steps", "app"])
        else:
            code_roots = _existing_roots(root, ["tests", "src", "features", "pages", "steps", "app"])
    if not test_data_roots:
        test_data_roots = _existing_roots(
            root,
            ["src/test_data", "test_data", "tests/data", "data", "resources", "tests/resources"],
        )

    feature_files = []
    for feature_root in feature_roots:
        feature_files.extend(_iter_files(feature_root, ["*.feature"]))

    code_patterns = ["*.java"] if framework_type == "java" else ["*.py"] if framework_type == "python" else ["*.java", "*.py"]
    code_files_raw = []
    for code_root in code_roots:
        code_files_raw.extend(_iter_files(code_root, code_patterns))

    workbook_files = [
        path for test_data_root in test_data_roots
        for path in _iter_files(test_data_root, ["*.xlsx", "*.xlsm"])
        if not path.name.startswith("~$")
    ]

    features = [_parse_feature(path, root) for path in feature_files]
    code_files = [_parse_code(path, root) for path in sorted(set(code_files_raw))]

    if lob_filter:
        features = [item for item in features if item.get("lob") == lob_filter]
        code_files = [
            item for item in code_files
            if item.get("lob") in {lob_filter, "", "general"} or item["step_definitions"]
        ]

    workbooks = []
    for path in workbook_files:
        try:
            workbooks.append(_scan_workbook(path, root, args.max_rows_per_sheet))
        except Exception as exc:
            workbooks.append({"path": _rel(path, root), "error": str(exc)})

    lobs = _build_lob_summary(features, code_files)
    return {
        "schema_version": 1,
        "generated_on": date.today().isoformat(),
        "source_root": str(root),
        "scope": {
            "framework_type": framework_type,
            "lob_filter": lob_filter or "all",
            "feature_roots": [_rel(path, root) for path in feature_roots],
            "code_roots": [_rel(path, root) for path in code_roots],
            "test_data_roots": [_rel(path, root) for path in test_data_roots],
            "max_rows_per_sheet": args.max_rows_per_sheet,
        },
        "summary": {
            "lob_count": len(lobs),
            "feature_file_count": len(features),
            "scenario_count": sum(item["scenario_count"] for item in features),
            "code_file_count": len(code_files),
            "java_file_count": len([item for item in code_files if item.get("language") == "java"]),
            "python_file_count": len([item for item in code_files if item.get("language") == "python"]),
            "workbook_count": len(workbooks),
        },
        "lobs": lobs,
        "features": features,
        "code_files": code_files,
        "java_files": [item for item in code_files if item.get("language") == "java"],
        "python_files": [item for item in code_files if item.get("language") == "python"],
        "workbooks": workbooks,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Scan a Java, Python, or mixed test framework and build a JSON knowledge file for MCP tools."
    )
    parser.add_argument("--root", default=str(PROJECT_ROOT), help="Framework root. Defaults to this repo.")
    parser.add_argument("--framework-type", default="auto", choices=["auto", "java", "python", "generic"])
    parser.add_argument("--lob", default="", help="Optional LOB key, e.g. auto, cyber, general-liability.")
    parser.add_argument("--features", default="", help="Feature folder relative to --root. Auto-detected when omitted.")
    parser.add_argument("--java", default="", help="Legacy Java code folder relative to --root. Prefer --code.")
    parser.add_argument("--code", default="", help="Code/test folder relative to --root. Auto-detected when omitted.")
    parser.add_argument("--test-data", default="", help="Excel test-data folder relative to --root. Auto-detected when omitted.")
    parser.add_argument("--output", default="docs/framework_knowledge.json", help="Output JSON path.")
    parser.add_argument("--max-rows-per-sheet", type=int, default=250, help="Limit sampled rows per Excel sheet.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    output = Path(args.output)
    if not output.is_absolute():
        output = root / output

    try:
        knowledge = build_knowledge(args)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(knowledge, indent=2), encoding="utf-8")
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    print(f"Wrote framework knowledge: {_rel(output, root)}")
    print(json.dumps(knowledge["summary"], indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
